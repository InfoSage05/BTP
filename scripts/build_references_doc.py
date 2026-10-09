# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt

entries = [
("Abbasian", """F. Abbasian, G.I. Hadaller, R.A. Fortman, J. Snell, S. Park,
An Artificial Neural Network (ANN) Model to Predict Critical Heat Flux (CHF) in a CANDU Fuel Element Simulation (FES) with Various Nonuniform Axial Heat Flux Shapes and Flow Liner Creep Profiles,
Nuclear Engineering and Design,
Volume 432,
2025,
113736,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2024.113736.
(https://www.sciencedirect.com/science/article/pii/S0029549324008367)"""),

("Abusah", """Johnson Lorlornyo Abusah, Shouxu Qiao, Ala Ayodeji, Samuel Yao Dedzie, Jiaxing Ren, Ruifeng Tian, Sichao Tan,
Deep-learning-based 3D bubble parameter prediction in a rod bundle subchannel: A single-view inference approach trained on dual-view reconstructions,
Nuclear Engineering and Design,
Volume 456,
2026,
115001,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2026.115001.
(https://www.sciencedirect.com/science/article/pii/S0029549326002530)"""),

("Abulawi", """Zaid Abulawi, Doyeong Lim, Abhiram Garimidi, Yang Liu,
Bayesian-optimized, feature-augmented deep ensemble for physics-guided critical heat-flux prediction with uncertainty quantification,
Annals of Nuclear Energy,
Volume 230,
2026,
112139,
ISSN 0306-4549,
https://doi.org/10.1016/j.anucene.2026.112139.
(https://www.sciencedirect.com/science/article/pii/S0306454926000277)"""),

("Ahmed", """Ibrahim Ahmed, Irene Gatti, Enrico Zio,
Optimized ensemble of neural networks for the prediction of critical heat flux,
Nuclear Engineering and Design,
Volume 439,
2025,
114111,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2025.114111.
(https://www.sciencedirect.com/science/article/pii/S0029549325002882)"""),

("Al Mahmud", """Abdullah Al Mahmud, Koji Morita, Wei Liu,
Critical heat flux prediction: A new approach adapting multi physics-aided machine learning,
International Communications in Heat and Mass Transfer,
Volume 173,
2026,
110879,
ISSN 0735-1933,
https://doi.org/10.1016/j.icheatmasstransfer.2026.110879.
(https://www.sciencedirect.com/science/article/pii/S0735193326004008)"""),

("Furlong", """Aidan Furlong, Xingang Zhao, Robert K. Salko, Xu Wu,
Physics-based hybrid machine learning for critical heat flux prediction with uncertainty quantification,
Applied Thermal Engineering,
Volume 279, Part A,
2025,
127447,
ISSN 1359-4311,
https://doi.org/10.1016/j.applthermaleng.2025.127447.
(https://www.sciencedirect.com/science/article/pii/S1359431125020393)"""),

("Groeneveld", """D.C. Groeneveld, J.Q. Shan, A.Z. Vasi\u0107, L.K.H. Leung, A. Durmayaz, J. Yang, S.C. Cheng, A. Tanase,
The 2006 CHF look-up table,
Nuclear Engineering and Design,
Volume 237,
2007,
Pages 1909-1922,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2007.02.014."""),

("Hall", """David D. Hall, Issam Mudawar,
Critical heat flux (CHF) for water flow in tubes\u2014I. Compilation and assessment of world CHF data,
International Journal of Heat and Mass Transfer,
Volume 43,
2000,
Pages 2573-2604,
ISSN 0017-9310,
https://doi.org/10.1016/S0017-9310(99)00191-X."""),

("Hardik1", """B.K. Hardik, S.V. Prabhu,
Boiling pressure drop, local heat transfer distribution and critical heat flux in helical coils with R123,
International Journal of Thermal Sciences,
Volume 125,
2018,
Pages 149-165,
ISSN 1290-0729,
https://doi.org/10.1016/j.ijthermalsci.2017.11.026.
(https://www.sciencedirect.com/science/article/pii/S1290072917306968)"""),

("Hardik2", """B.K. Hardik, Gajendra Kumar, S.V. Prabhu,
Boiling pressure drop, local heat transfer distribution and critical heat flux in horizontal straight tubes,
International Journal of Heat and Mass Transfer,
Volume 113,
2017,
Pages 466-481,
ISSN 0017-9310,
https://doi.org/10.1016/j.ijheatmasstransfer.2017.05.073.
(https://www.sciencedirect.com/science/article/pii/S0017931016340443)"""),

("Hardik3", """B.K. Hardik, S.V. Prabhu,
Critical heat flux in helical coils at low pressure,
Applied Thermal Engineering,
Volume 112,
2017,
Pages 1223-1239,
ISSN 1359-4311,
https://doi.org/10.1016/j.applthermaleng.2016.10.114.
(https://www.sciencedirect.com/science/article/pii/S1359431116325650)"""),

("Huang", """Zili Huang, Yihua Duo, Hong Xu,
Prediction of two-phase flow patterns based on machine learning,
Nuclear Engineering and Design,
Volume 421,
2024,
113107,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2024.113107.
(https://www.sciencedirect.com/science/article/pii/S0029549324002073)"""),

("Kefer", """V. Kefer, W. K\u00f6hler, W. Kastner,
Critical heat flux (CHF) and post-CHF heat transfer in horizontal and inclined evaporator tubes,
International Journal of Multiphase Flow,
Volume 15, Issue 3,
1989,
Pages 385-392,
ISSN 0301-9322,
https://doi.org/10.1016/0301-9322(89)90008-6.
(https://www.sciencedirect.com/science/article/pii/0301932289900086)"""),

("Khalid1", """Rehan Zubair Khalid, Atif Khan, Farooq Ahmad, Muthanna H. Al-Dahhan, Atta Ullah,
The effect of reduction of input-parameters on data-driven prediction of critical heat flux in rod bundles using extended parameters for reactor thermal\u2013hydraulic safety,
International Communications in Heat and Mass Transfer,
Volume 172, Part 6,
2026,
110659,
ISSN 0735-1933,
https://doi.org/10.1016/j.icheatmasstransfer.2026.110659.
(https://www.sciencedirect.com/science/article/pii/S0735193326001806)"""),

("Khalid2", """Rehan Zubair Khalid, Ibrahim Ahmed, Atta Ullah, Enrico Zio, Asifullah Khan,
Enhancing accuracy of prediction of critical heat flux in Circular channels by ensemble of deep sparse autoencoders and deep neural Networks,
Nuclear Engineering and Design,
Volume 429,
2024,
113587,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2024.113587.
(https://www.sciencedirect.com/science/article/pii/S0029549324006873)"""),

("Liu", """Haidong Liu, Hongfu Jiang, Deqi Chen, Jiang Qin, Yangyang Wang, Mingxia Liu,
Critical heat flux prediction for annular channel through application of machine learning techniques,
International Communications in Heat and Mass Transfer,
Volume 167, Part B,
2025,
109279,
ISSN 0735-1933,
https://doi.org/10.1016/j.icheatmasstransfer.2025.109279.
(https://www.sciencedirect.com/science/article/pii/S0735193325007055)"""),

("Marcinkiewicz", """Krzysztof Marcinkiewicz, Ole Wieckhorst, Rafael Maci\u00e1n-Juan, Markus Rehm, Peter Pohl,
Recurrent neural network-based prediction of critical heat flux in rod bundles with non-uniform axial power shape,
Nuclear Engineering and Design,
Volume 395,
2022,
111825,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2022.111825.
(https://www.sciencedirect.com/science/article/pii/S0029549322001790)"""),

("Rashidi", """M.M. Rashidi, M. Alhuyi Nazari, C. Harley, E. Momoniat, I. Mahariq, N. Ali,
Applications of machine learning methods for boiling modeling and prediction: A comprehensive review,
Chemical Thermodynamics and Thermal Analysis,
Volume 8,
2022,
100081,
ISSN 2667-3126,
https://doi.org/10.1016/j.ctta.2022.100081."""),

("Serrao", """Bruno P. Serrao, Ye Kwon Huh, Eliot Ciuperca, Elvan Sahin, Kaibo Liu, Juliana P. Duarte,
A quantitative analysis of ATF surface characteristics on critical heat flux using Machine learning,
Nuclear Engineering and Design,
Volume 435,
2025,
113924,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2025.113924.
(https://www.sciencedirect.com/science/article/pii/S0029549325001013)"""),

("Wu", """Huakang Wu, Minyang Gui, Di Wu,
Physics-Informed hybrid machine learning for critical heat flux prediction: A comparative analysis of modeling approaches,
Nuclear Engineering and Design,
Volume 445,
2025,
114434,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2025.114434.
(https://www.sciencedirect.com/science/article/pii/S0029549325006119)"""),

("Yang", """Han Yang, Dengke Li, Zhengkai Wang, Yucheng Yan, Feng Luan,
LUT-guided grey-box neural networks for predicting critical heat flux in uniformly heated water-cooled circular tubes,
Nuclear Engineering and Design,
Volume 458,
2026,
115112,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2026.115112.
(https://www.sciencedirect.com/science/article/pii/S002954932600364X)"""),

("Zhang", """Wei Zhang, Junsen Fu, Shuo Chen, Lijun Yu, Yao Xiao, Hanyang Gu,
Investigation of a hybrid neural network framework for CHF prediction in a wire-wrapped rod bundle,
Annals of Nuclear Energy,
Volume 226,
2026,
111804,
ISSN 0306-4549,
https://doi.org/10.1016/j.anucene.2025.111804.
(https://www.sciencedirect.com/science/article/pii/S0306454925006218)"""),

("Zhao", """Xingang Zhao, Robert K. Salko, Koroush Shirvan,
Improved departure from nucleate boiling prediction in rod bundles using a physics-informed machine learning-aided framework,
Nuclear Engineering and Design,
Volume 374,
2021,
111084,
ISSN 0029-5493,
https://doi.org/10.1016/j.nucengdes.2021.111084.
(https://www.sciencedirect.com/science/article/pii/S0029549321000364)"""),
]

doc = docx.Document()

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

title = doc.add_heading('References', level=1)

for i, (_, text) in enumerate(entries, start=1):
    lines = text.strip().split('\n')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(f"[{i}] ")
    run.bold = True
    for j, line in enumerate(lines):
        if j > 0:
            p.add_run('\n')
        p.add_run(line)

out_path = r"D:\BTP\CHF_References.docx"
doc.save(out_path)
print("Saved:", out_path, "entries:", len(entries))
