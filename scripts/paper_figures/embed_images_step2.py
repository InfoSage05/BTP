"""
Step 2: insert the three new tree-model parity plots (Random Forest,
Gradient Boosting, XGBoost -- requested at the Sept 17, 2026 team meeting)
as Figures 16-18, directly after Figure 15 and before Section 3.7.
"""
from docx import Document
from docx.shared import Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

SRC = "CHF_Research_Paper_step1.docx"
DST = "CHF_Research_Paper.docx"

NEW_FIGURES = [
    (16, "docs/manuscript/fig_parity_plot_random_forest.png",
     "Figure 16. Predicted vs. measured CHF for a Random Forest model, on the same real "
     "NRC interpolation and extrapolation test splits used for Figure 15, so the two models "
     "are directly comparable. Both axes are log scaled; the dashed lines mark a +/-10% "
     "error band around perfect prediction."),
    (17, "docs/manuscript/fig_parity_plot_gradient_boosting.png",
     "Figure 17. Predicted vs. measured CHF for a Gradient Boosting model, on the same real "
     "NRC interpolation and extrapolation test splits used for Figures 15 and 16."),
    (18, "docs/manuscript/fig_parity_plot_xgboost.png",
     "Figure 18. Predicted vs. measured CHF for an XGBoost model, on the same real NRC "
     "interpolation and extrapolation test splits used for Figures 15-17."),
]

DISPLAY_WIDTH_EMU = 5760720  # same displayed width as Figure 15

doc = Document(SRC)

anchor = None
for p in doc.paragraphs:
    if p.text.strip() == "3.7 Explainable AI":
        anchor = p
        break
assert anchor is not None, "could not find Section 3.7 heading"

intro = anchor.insert_paragraph_before(
    "Figures 16-18 extend the same parity-plot comparison to the other three models "
    "evaluated in this paper -- Random Forest, Gradient Boosting, and XGBoost -- on the "
    "identical real NRC interpolation and extrapolation test splits used for Figure 15, "
    "so all four models can be read side by side under the same axes and the same error band."
)

for fig_num, img_path, caption in NEW_FIGURES:
    im = Image.open(img_path)
    ratio = im.width / im.height
    height_emu = round(DISPLAY_WIDTH_EMU / ratio)

    spacer = anchor.insert_paragraph_before("")

    pic_p = anchor.insert_paragraph_before()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_p.add_run()
    run.add_picture(img_path, width=Emu(DISPLAY_WIDTH_EMU), height=Emu(height_emu))

    cap_p = anchor.insert_paragraph_before()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap_p.add_run(caption)
    cap_run.italic = True
    cap_run.font.size = Pt(9)

doc.save(DST)
print(f"wrote {DST}")
